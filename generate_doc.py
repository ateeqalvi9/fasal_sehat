import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    PRIMARY_COLOR = RGBColor(0x2E, 0x7D, 0x32) # #2E7D32 Forest Green
    SECONDARY_COLOR = RGBColor(0x1B, 0x5E, 0x20)
    TEXT_COLOR = RGBColor(0x1B, 0x1B, 0x1B)
    MUTED_COLOR = RGBColor(0x55, 0x55, 0x55)

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run('Fasal Sehat (فصل صحت)')
    title_run.font.name = 'Segoe UI'
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = PRIMARY_COLOR

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(16)
    sub_run = subtitle_p.add_run('Comprehensive Technical & Academic Guide: Questions & Answers\nPrepared for Mobile Computing & AI Course Projects')
    sub_run.font.name = 'Segoe UI'
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = MUTED_COLOR

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = SECONDARY_COLOR
        return p

    def add_p(text, bold_prefix='', italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = 'Segoe UI'
            r_b.font.size = Pt(10.5)
            r_b.font.bold = True
            r_b.font.color.rgb = TEXT_COLOR
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10.5)
        r.font.italic = italic
        r.font.color.rgb = TEXT_COLOR
        return p

    def add_bullet(text, bold_prefix=''):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = 'Segoe UI'
            r_b.font.size = Pt(10.5)
            r_b.font.bold = True
            r_b.font.color.rgb = TEXT_COLOR
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10.5)
        r.font.color.rgb = TEXT_COLOR
        return p

    def set_cell_background(cell, hex_color):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    # Q1
    add_heading_1('Q1. Do we save data in a database?')
    add_p('Yes, but exclusively on the local Android device. Because "Fasal Sehat" is designed as a strict offline-first application for farmers in rural agricultural fields with no cellular reception or Wi-Fi, there is NO cloud database (no Firebase, MongoDB, or AWS) and zero network permissions.')

    add_heading_2('1.1 What goes into the SQLite Database (sqflite):')
    add_bullet(' Created automatically on the device when a user diagnoses a leaf. It records: scan ID, disease class label, confidence score (e.g. 0.88), severity tag (healthy, mild, severe), file path to the compressed image, file path to the 96x96 thumbnail, and the timestamp.', bold_prefix='Diagnosis History Table (`history`):')
    add_bullet(' Raw photo bytes (blobs) are NEVER stored in SQLite. Instead, photos are compressed to JPEG (max 800px, quality 80) and stored in the app\'s internal documents directory, with only the local file paths stored in SQLite. This prevents memory bloat and keeps the SQLite database under 1 MB.', bold_prefix='Storage Optimization:')
    add_bullet(' The SQLite history is capped at 200 entries (FIFO — deleting the oldest entry automatically) to bound storage consumption on budget 2 GB RAM / 16 GB ROM phones.', bold_prefix='200-Item FIFO Cap:')

    add_heading_2('1.2 What is bundled as Static Assets (Read-Only JSON):')
    add_bullet(' Stored as a local JSON file (`assets/data/remedies.json`) bundled directly inside the APK. It maps each disease class to its localized English and Urdu names, severity tier, 2–3 actionable treatment steps with icons, and audio clip paths. It is read once at app launch and cached in memory.', bold_prefix='Remedy Knowledge Base:')

    add_heading_2('1.3 What does NOT go into the mobile database:')
    add_bullet(' The thousands of raw training photos stay solely in Google Colab / Kaggle. Only the final compiled, quantized model (`crop_model.tflite` <= 5 MB) is placed inside `assets/model/` inside the Flutter project.', bold_prefix='Training Images:')
    add_bullet(' Zero remote servers, user profiles, credentials, or tracking tables exist in this app.', bold_prefix='No User Accounts or Server Tables:')

    # Q2
    add_heading_1('Q2. Do we need to train a model, or is everything already available?')
    add_p('You do not train from scratch, but you MUST fine-tune a model. This is standard Transfer Learning and forms the core of your AI course evaluation.')

    add_heading_2('2.1 Component Breakdown & Responsibility Table:')
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Component', 'Already Available?', 'Your Exact Role', 'Course Alignment']
    widths = [Inches(1.5), Inches(1.2), Inches(2.4), Inches(1.5)]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.name = 'Segoe UI'
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr_cells[i], '2E7D32')
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)

    rows_data = [
        ('MobileNetV2 Architecture & Pretrained ImageNet Weights', 'Yes (Pretrained by Google)', 'Download via tf.keras.applications.MobileNetV2(weights="imagenet", include_top=False).', 'AI Course (Base backbone)'),
        ('Crop Disease Classification Head', 'No', 'Fine-tune in Google Colab on your chosen crop dataset (2–3 hours). Freeze base layers, train top classifier, fine-tune last 20 layers, evaluate with confusion matrix, and quantize to .tflite.', 'AI Course (Primary ML Deliverable)'),
        ('Flutter Offline App & Inference Engine', 'No', 'Build the offline Flutter app with camera preview, background isolate image processing, TFLite interpreter, SQLite history, and audio playback.', 'Mobile Computing Course (Primary App Deliverable)')
    ]

    for row_idx, data in enumerate(rows_data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = 'F5F5F5' if row_idx % 2 == 1 else 'FFFFFF'
        for col_idx, text in enumerate(data):
            row_cells[col_idx].text = text
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.name = 'Segoe UI'
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = TEXT_COLOR
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)

    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

    add_heading_2('2.2 Why downloading a pre-existing disease model is risky:')
    add_p('While ready-made disease models exist on GitHub or Kaggle, simply downloading a pre-trained `.tflite` makes your AI course submission superficial and vulnerable to academic scrutiny. By fine-tuning MobileNetV2 yourself in Colab, you obtain authentic training logs, loss curves, confusion matrices, and precision/recall metrics for your viva and report. You can keep a downloaded model only as an emergency backup.')

    # Q3
    add_heading_1('Q3. Where is the data? Which diseases should we target?')
    add_p('The dataset provides both the images and the class names simultaneously. In image classification datasets, folder names ARE the disease labels (e.g., `Tomato___Early_blight`, `Potato___Late_blight`, `Tomato___healthy`).')

    add_heading_2('3.1 Primary Data Sources:')
    add_bullet(' The gold-standard starting point containing 54,305 lab-condition images across 38 crop classes. Hosted on Kaggle: https://www.kaggle.com/datasets/abdallahalidev/plantvillage-dataset or https://www.kaggle.com/datasets/emmarex/plantdisease/data', bold_prefix='1. PlantVillage Dataset (Kaggle):')
    add_bullet(' Real-world field photos under natural sunlight, varying backgrounds, and shadows. Hosted on GitHub: https://github.com/pratikkayal/PlantDoc-Object-Detection-Dataset. Cropped images from PlantDoc bridge the domain gap between lab photos and actual farmer fields.', bold_prefix='2. PlantDoc Dataset (Field Conditions):')
    add_bullet(' PlantVillage covers Tomato, Potato, Pepper, and Corn. If you choose major Pakistani cash crops (Wheat, Cotton, Rice, Sugarcane), use Mendeley Data or Kaggle (e.g. "Cotton Leaf Disease Dataset" or "Rice Leaf Disease Dataset").', bold_prefix='3. Pakistani Regional Crops (Mendeley Data):')

    add_heading_2('3.2 Recommended Crop Selection (Balanced Scope):')
    add_p('For a university semester project, select 3 to 4 crops with 8 to 14 classes total, plus one negative class:')
    add_bullet('Early Blight, Late Blight, Yellow Leaf Curl Virus, Healthy', bold_prefix='Tomato (4 classes): ')
    add_bullet('Early Blight, Late Blight, Healthy', bold_prefix='Potato (3 classes): ')
    add_bullet('Bacterial Spot, Healthy', bold_prefix='Pepper / Bell Pepper (2 classes): ')
    add_bullet('Crucial 500+ random non-leaf pictures (hands, shoes, grass, dirt, clothes). This ensures that if a farmer points the phone at anything other than a crop leaf, the app accurately displays "Could not identify / Retake photo" instead of outputting a bogus disease.', bold_prefix='Negative Class (`not_a_leaf`): ')

    # Q4
    add_heading_1('Q4. Where do the cures come from?')
    add_p('There is no automated download that provides "disease plus remedy." You construct the remedy knowledge base manually by referencing vetted agronomy databases.')

    add_heading_2('4.1 Authoritative Verification Sources:')
    add_bullet(' Managed by CABI, providing over 6,000 verified factsheets in multiple languages covering symptoms, chemical controls, and cultural management: https://www.plantwise.org/knowledgebank/', bold_prefix='CABI Plantwise Knowledge Bank: ')
    add_bullet(' Provincial Agriculture Extension Departments (Punjab, Sindh, KP) publish farmer advisory guides and seasonal spray schedules.', bold_prefix='PARC & Provincial Extension Services: ')
    add_bullet(' Research publications and extension bulletins for specific crop pathology.', bold_prefix='University of Agriculture Faisalabad (UAF): ')

    add_heading_2('4.2 Authoring Rules for Farmer Safety & Practicality:')
    add_bullet('Provide 2 to 3 concise, actionable steps paired with icons (e.g. 💧 Watering, ✂️ Pruning, 🧴 Fungicide spray).', bold_prefix='1. Icon-First & Short: ')
    add_bullet('Never specify commercial brand names (which vary by region) or exact concentrated chemical dosages (which carry liability of crop chemical burns). Prescribe generic classes such as "copper-based fungicide" or "neem oil extract".', bold_prefix='2. Generic Recommendations Only: ')
    add_bullet('For any spray recommendation, include a mandatory step: "Wear protective gloves and mask when spraying."', bold_prefix='3. Safety Gear Mandate: ')
    add_bullet('Every result card must display the fixed disclaimer: "For severe cases consult your local agriculture extension officer."', bold_prefix='4. Mandatory Advisory Disclaimer: ')
    add_bullet('Translate remedies into simple, colloquial Urdu and record 15-second clear voice audio clips for low-literacy farmers.', bold_prefix='5. Urdu Audio Pairing: ')

    # Order of Work
    add_heading_1('Your Recommended Step-by-Step Order of Work')
    steps = [
        ('Step 1: Crop & Class Selection —', 'Select 3–4 crops (e.g. Tomato, Potato, Pepper) totaling 8–14 classes plus `not_a_leaf`.'),
        ('Step 2: Dataset Preparation —', 'Download subsets from PlantVillage and blend with 500+ non-leaf images in Colab.'),
        ('Step 3: Transfer Learning in Google Colab —', 'Load MobileNetV2, train top classification head, fine-tune top 20 layers, evaluate with confusion matrix & F1 scores.'),
        ('Step 4: Export & Quantization —', 'Convert Keras model to float16 or dynamic-range quantized `.tflite` (target size <= 5 MB).'),
        ('Step 5: Remedy Knowledge Base Authoring —', 'Compile `assets/data/remedies.json` with English and Urdu steps, cross-verified against CABI Plantwise.'),
        ('Step 6: Mobile App Implementation (Flutter) —', 'Complete Phase 1 (UI), Phase 2 (Camera), Phase 3 (TFLite ML pipeline), and Phase 4 (SQLite history).'),
        ('Step 7: Urdu Audio Recording —', 'Record short 10–15 second `.opus` audio clips explaining each remedy in spoken Urdu.'),
        ('Step 8: Verification & Defense Prep —', 'Test APK in Airplane Mode on an Android phone, document latency (<1.5s) and APK size (<25MB), and prepare project report.')
    ]

    for title, desc in steps:
        add_bullet(desc, bold_prefix=title)

    # Save documents
    output_path1 = r'C:\Users\gamers arena\Downloads\Fasal_Sehat_Questions_and_Answers.docx'
    output_path2 = r'C:\Mobile_Computing\fasal_sehat\Fasal_Sehat_Questions_and_Answers.docx'

    doc.save(output_path1)
    doc.save(output_path2)
    print(f'Successfully generated DOCX files at:\n{output_path1}\n{output_path2}')

if __name__ == '__main__':
    create_document()
